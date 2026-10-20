from openai import OpenAI
import json
import os
from pypdf import PdfReader
import pyalex
from pyalex import Works
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from functools import wraps
import time
import io
import pandas as pd

class LitReview:
    citations = None
    compiled_citations = None
    review = None
    
    def __init__(self, api_key, email):
        os.environ['OPENAI_API_KEY'] = api_key
        
        self.client = OpenAI(
            base_url = "https://openrouter.ai/api/v1"
        )

    def get_completion(self, system_prompt, prompt, json_mode = False, json_schema = None):
        if json_schema:
            system_prompt = system_prompt + f"""\n<json_template>\n{json_schema}\n</json_template>"""
        kwargs = {
            "model": "openrouter/free",
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": prompt}
            ]
        }
        if json_mode:
            kwargs["response_format"] = {"type": "json_schema",
                                         "json_schema": json_schema}
    
        completion = self.client.chat.completions.create(**kwargs)
        return completion.choices[0].message.content

    def get_topics(self, uploaded_file):
        reader = PdfReader(uploaded_file)
        content = ""
        for index, page in enumerate(reader.pages):
            content = content + page.extract_text()
    
        system_prompt = f"""
    <instructions>
    You are an academic research assistant tasked with developing a review of literature. 
    Read through the following document and compile a list of academic/research topics related to the described program content. 
    Make sure to include important specifics such as region or social group, if applicable, within the topics.
    Optimize the topics for a semantic search API for academic literature, combining optimal length of topic description and keywords.
    Return a json object with a list of topics related to the program described.
    Strictly follow the exact json schema described. 
    Return only the populated JSON without any explanations, reasoning or decorators.
    </instructions>
        """
    
        json_template = {'name': 'topics_schema',
                         'strict': True,
                         'schema': {
                             'type': 'object',
                             'properties': {
                                 'topics': {'type': 'array',
                                            'items': {
                                                'type': 'string'
                                                }
                                            }
                                        }
                                }
                        }
    
        topics = self.get_completion(system_prompt = system_prompt,
                                prompt = content,
                                json_mode = True,
                                json_schema = json_template)
    
        topics = topics.strip('`')
        topics = topics.strip('json')

        return topics

    def get_citations(self, topic, num_papers):    
        return (
            Works()
            .similar(topic)
            .get(per_page = num_papers)
            )

    def compile_citations(self, citations):
        compiled_citations = dict()
        for work in citations:
            authors = work.get('authorships')
            authors = [author.get('author').get('display_name') for author in authors]
            year = work.get('publication_year')
            text = work.get('abstract_inverted_index')
                
            if len(authors) == 0 or not text: continue
            if len(authors) == 1:
                short_ref = authors[0].split(" ")[-1] + ', ' + str(year)
            elif len(authors) == 2:
                short_ref = authors[0].split(" ")[-1] + ' & ' + authors[1].split(" ")[-1] + ', ' + str(year)
            elif len(authors) == 3:
                short_ref = authors[0].split(" ")[-1] + ', ' + authors[1].split(" ")[-1] + ' & ' + authors[2].split(" ")[-1] + ', ' + str(year)
            else:
                short_ref = authors[0].split(" ")[-1] + ' et. al., ' + str(year)
                
            max_index = max(max(positions) for positions in text.values())
            word_list = [""]*(max_index + 1)
                
            for word, positions in text.items():
                for pos in positions:
                    word_list[pos] = word
    
            abstract = " ".join(word_list)
            title = work.get('title')
            source = work.get('primary_location').get('raw_source_name')
            doi = work.get('doi')
            author_list = [author.split(" ")[-1] + ", " + ". ".join([a[0] for a in author.split(" ") if a][:-1]) + "." for author in authors]
            year_style = '(' + str(year) + ') ' if year else '(n.d.)'
            doi_style = ", doi: " + doi if doi else ""
            source_style = source if source else ""
            volume = work.get('biblio').get('volume')
            issue = work.get('biblio').get('issue')
            voliss_style = str(volume) + '(' + str(issue) + ')' if volume and issue else ""
            apa_ref = ", ".join(author_list) + year_style + ' "' + title + '", ' + source_style + " " + voliss_style + doi_style
    
            compiled_citations[short_ref] = {"short_ref": short_ref, "Year": year, "abstract": abstract, "apa": apa_ref, "doi": doi}
        return compiled_citations

    def lit_review(self, topic, papers):
        system_prompt = f"""
    <instructions>
    You are a high level academic research assistant. 
    Your job is to take the paper abstracts provided below and produce a short review of literature on the topic given.
    Limit the output to between {len(papers) * 7} and {len(papers) * 39} words.
    Use only the paper abstracts provided below and cite them in the text as per the style outlined in academic format.
    The input is in the form: <in-text citation>Abstract</in-text citation> for each paper abstract included.
    Make sure that the output text matches the academic format specified in <academic format> and do not use any different format.
    </instructions>
    
    <academic format>
    Lorem ipsum dolor sic amet (Author 1 & Author 2, Year), consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua (Author, Year).
    Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat (Author et. al., Year).
    </academic format>
        """
        prompt = f"""
    <topic>
    {topic}
    \n\n
    </topic>
    
    <related papers>
    {"\n".join(['<' + authors + '>' + ': ' + content['abstract'] + '</' + authors + '>' for authors, content in papers.items()])}
    </related papers>
    """
        lit_review = ""
        while len(lit_review) < 30:
            lit_review = self.get_completion(system_prompt, prompt, json_mode = False)
        return lit_review

    def excel_export(self, compiled_citations):
        papers = []

        for topic, citations in compiled_citations.items():
            for paper_id, details in citations.items():
                row = details.copy()
                row["topic"] = topic
                papers.append(row)

        df = pd.DataFrame(papers)

        # Put topic first
        if not df.empty:
            cols = ["topic"] + [c for c in df.columns if c != "topic"]
            df = df[cols]

        buffer = io.BytesIO()

        with pd.ExcelWriter(buffer, engine="xlsxwriter") as writer:
            df.to_excel(writer, sheet_name="papers", index=False)

            worksheet = writer.sheets["papers"]
            worksheet.freeze_panes(1, 0)

            if len(df.columns) > 0:
                worksheet.autofilter(0, 0, len(df), len(df.columns) - 1)

        return buffer.getvalue()
        
        

    def word_export(self, literature_review, compiled_citations):
        buffer = io.BytesIO()
        doc = Document()
        
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Lato'
        font.size = Pt(11)
        
        h0_style = doc.styles['Title']
        h0_style.font.name = 'Lato'
        h0_style.size = Pt(20)
        h0_style.font.color.rgb = RGBColor(126, 63, 152)
        
        h1_style = doc.styles['Heading 1']
        h1_style.font.name = 'Lato'
        h1_style.size = Pt(16)
        h1_style.font.color.rgb = RGBColor(0, 173, 141)
        
        h2_style = doc.styles['Heading 2']
        h2_style.font.name = 'Lato'
        h2_style.size = Pt(13)
        h2_style.font.color.rgb = RGBColor(191, 159, 203)
    
        # Access the first section of the document
        section = doc.sections[0]
        
        # Set margins to 1 inch
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        h = doc.add_heading("Literature Review", level = 0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        for topic, review in literature_review.items():
            h = doc.add_heading(topic, level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            p = doc.add_paragraph()
            p.add_run(review)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        sorted_papers = []

        for topic, citations in compiled_citations.items():
            for authors, details in citations.items():
                sorted_papers.append(details['apa'])
    
        sorted_papers.sort()

        doc.add_heading("References", level = 0)
        for paper in sorted_papers:
            p = doc.add_paragraph()
            p.add_run(paper)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = -Inches(0.5)
        
        
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def generate_review(self):
        if not self.topics: self.get_topics()
    
        if not self.citations: self.get_citations()
    
        if not self.compiled_citations: self.compile_citations()
    
        if not self.review: self.lit_review()
    

    def export_review(self, out_path):
        word_export(self, self.review, self.compiled_citations, out_path)
